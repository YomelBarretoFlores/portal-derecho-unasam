from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = (
    ROOT
    / "docs"
    / "infrastructure"
    / "Informe_tecnico_solicitud_alojamiento_SSH_APA7_Derecho_UNASAM.docx"
)
SKILL_SCRIPTS = Path(
    "/Users/yomeljairbarretoflores/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.805.11740/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry  # noqa: E402


FONT = "Times New Roman"
BLACK = "000000"
GRAY = "777777"
PAGE_WIDTH_DXA = 9360


def set_run(run, *, size=12, bold=False, italic=False):
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), FONT)
    r_pr.rFonts.set(qn("w:hAnsi"), FONT)
    r_pr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = None


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_border(cell, *, top=None, bottom=None, left=None, right=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    definitions = {"top": top, "bottom": bottom, "left": left, "right": right}
    for edge, spec in definitions.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        if spec is None:
            node.set(qn("w:val"), "nil")
        else:
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(spec.get("size", 8)))
            node.set(qn("w:color"), spec.get("color", BLACK))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run(run, size=12)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    add_page_number(section.header.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.italic = style_name == "Heading 3"
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def add_body(doc, text, *, indent=True, bold_lead=None, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = keep
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    if bold_lead and text.startswith(bold_lead):
        set_run(paragraph.add_run(bold_lead), bold=True)
        set_run(paragraph.add_run(text[len(bold_lead):]))
    else:
        set_run(paragraph.add_run(text))
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_run(paragraph.add_run(text))
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    set_run(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_run(paragraph.add_run(text), bold=True, italic=level == 3)
    return paragraph


def add_table_label(doc, number, title):
    number_p = doc.add_paragraph()
    number_p.paragraph_format.line_spacing = 2
    number_p.paragraph_format.keep_with_next = True
    set_run(number_p.add_run(f"Tabla {number}"), bold=True)
    title_p = doc.add_paragraph()
    title_p.paragraph_format.line_spacing = 2
    title_p.paragraph_format.keep_with_next = True
    set_run(title_p.add_run(title), italic=True)


def add_apa_table(doc, headers, rows, widths, *, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_repeat_header(table.rows[0])
    outer = {"size": 10, "color": BLACK}
    inner = {"size": 6, "color": GRAY}

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell, top=outer, bottom=inner)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        set_run(paragraph.add_run(header), size=font_size, bold=True)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        is_last = row_index == len(rows) - 1
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, bottom=outer if is_last else None)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            set_run(paragraph.add_run(str(value)), size=font_size)

    apply_table_geometry(
        table,
        widths,
        indent_dxa=120,
        cell_margins_dxa={"top": 90, "bottom": 90, "start": 120, "end": 120},
    )
    after = doc.add_paragraph()
    after.paragraph_format.line_spacing = 2
    return table


def add_title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 2
    title.paragraph_format.space_after = Pt(0)
    set_run(
        title.add_run(
            "Informe técnico para la solicitud de alojamiento institucional del Portal Web de la Facultad de Derecho y Ciencias Políticas"
        ),
        bold=True,
    )
    for text in (
        "Facultad de Derecho y Ciencias Políticas",
        "Universidad Nacional Santiago Antúnez de Mayolo",
        "Documento dirigido al Área de Tecnologías de la Información",
        "Responsable técnico: ________________________________",
        "12 de agosto de 2026",
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.space_after = Pt(0)
        set_run(paragraph.add_run(text))
    doc.add_page_break()


def add_summary(doc):
    add_heading(doc, "Resumen", 1)
    add_body(
        doc,
        "El presente informe define los requerimientos mínimos que la Facultad de Derecho y Ciencias Políticas debe solicitar al Área de Tecnologías de la Información para publicar su portal web en la infraestructura de la Universidad. La modalidad propuesta puede consistir en la asignación de un espacio dentro de un servidor Linux institucional, con acceso mediante Secure Shell (SSH) restringido a la red de la UNASAM o a una red privada virtual (VPN). No es indispensable que la Facultad reciba una máquina virtual exclusiva, siempre que el espacio asignado permita ejecutar la aplicación, conectarse a PostgreSQL, conservar archivos de forma persistente y publicar el servicio bajo un subdominio institucional con HTTPS.",
    )
    add_body(
        doc,
        "Además del acceso al servidor, el área responsable debe proporcionar una base de datos PostgreSQL, una ubicación persistente para documentos e imágenes, la configuración del subdominio, el certificado digital y un procedimiento de respaldo. El informe incluye un modelo de los datos de acceso que podrían entregarse en un archivo de texto, de manera semejante al servidor institucional denominado anteriormente TESTER.",
    )
    keywords = doc.add_paragraph()
    keywords.paragraph_format.line_spacing = 2
    keywords.paragraph_format.first_line_indent = Inches(0.5)
    set_run(keywords.add_run("Palabras clave: "), italic=True)
    set_run(keywords.add_run("alojamiento institucional, SSH, Laravel, PostgreSQL, almacenamiento persistente, subdominio"))
    doc.add_page_break()


def add_main_content(doc):
    add_heading(doc, "Informe técnico para la solicitud de alojamiento institucional", 1)

    add_heading(doc, "Objeto del requerimiento", 2)
    add_body(
        doc,
        "El objeto del requerimiento es obtener un espacio de producción dentro de la infraestructura tecnológica de la UNASAM para alojar el Portal Web de la Facultad de Derecho y Ciencias Políticas. La solicitud no debe limitarse a pedir una cuenta SSH, porque dicha cuenta constituye únicamente el medio de acceso. El servicio completo debe comprender ejecución de la aplicación, conectividad con PostgreSQL, almacenamiento de archivos, publicación web, seguridad, copias de respaldo y soporte operativo.",
    )
    add_body(
        doc,
        "La solución puede implementarse en un servidor compartido, una máquina virtual o una plataforma institucional de contenedores. Para el modelo utilizado anteriormente por la Universidad, es suficiente que TI asigne un directorio o espacio aislado dentro de un servidor Linux y entregue los datos de conexión correspondientes, siempre que se cumplan los requisitos indicados en este documento.",
    )

    add_heading(doc, "Configuración técnica verificada del sistema", 2)
    add_body(
        doc,
        "Las versiones fueron verificadas directamente en los archivos instalados, el archivo de bloqueo de dependencias y la imagen de despliegue del proyecto. La instalación exacta requiere PHP 8.4.1 o una versión posterior compatible, debido a las dependencias actualmente bloqueadas.",
    )
    add_table_label(doc, 1, "Componentes principales verificados")
    add_apa_table(
        doc,
        ["Componente", "Versión o requisito"],
        [
            ("Laravel Framework", "13.12.0"),
            ("Filament", "5.6.6"),
            ("PHP", "8.4.1 como mínimo efectivo; imagen Docker basada en PHP 8.4"),
            ("Base de datos", "PostgreSQL"),
            ("Construcción de interfaz", "Node.js 22 y Vite durante compilación"),
            ("Servidor web", "FrankenPHP/Caddy, Nginx con PHP-FPM o Apache compatible"),
        ],
        [2600, 6760],
    )

    add_heading(doc, "Modalidad recomendada de acceso", 2)
    add_body(
        doc,
        "La modalidad más directa consiste en que TI entregue una dirección IP o nombre interno, puerto de conexión, usuario y mecanismo de autenticación. El acceso puede limitarse a equipos conectados a la red institucional. Cuando el responsable técnico trabaje fuera de la Universidad, TI puede exigir el uso de VPN.",
    )
    add_body(
        doc,
        "La autenticación mediante llave pública SSH es preferible porque permite revocar accesos individuales sin compartir contraseñas. Sin embargo, si el procedimiento institucional vigente utiliza usuario y contraseña, también puede utilizarse, siempre que la contraseña se entregue mediante un canal seguro, sea personal y pueda cambiarse. No se deben incorporar contraseñas en el oficio, repositorio de código o documentación pública.",
    )

    add_heading(doc, "Información que debe proporcionar el Área de TI", 2)
    add_body(doc, "El documento de entrega técnica debe contener, como mínimo, los siguientes datos:")
    for item in (
        "Nombre o identificador del servidor, por ejemplo, PRODUCCION-DERECHO o el nombre institucional aprobado.",
        "Dirección IP o nombre DNS interno del servidor.",
        "Puerto SSH y confirmación de acceso desde la red institucional o mediante VPN.",
        "Usuario asignado y mecanismo de autenticación: contraseña temporal o llave pública.",
        "Ruta absoluta destinada al proyecto y política de permisos o uso de sudo.",
        "Responsable de TI y canal de atención ante indisponibilidad o bloqueos de acceso.",
    ):
        add_bullet(doc, item)

    add_table_label(doc, 2, "Modelo de datos para la entrega del acceso al servidor")
    add_apa_table(
        doc,
        ["Dato", "Valor que debe completar TI"],
        [
            ("Servidor", "PRODUCCION-DERECHO, TESTER u otro nombre aprobado"),
            ("IP o hostname", "Dirección privada o nombre institucional"),
            ("Puerto SSH", "22 o puerto definido por TI"),
            ("Usuario", "Cuenta nominativa del responsable técnico"),
            ("Autenticación", "Llave SSH o contraseña temporal"),
            ("Acceso de red", "Red UNASAM, VPN o lista de direcciones autorizadas"),
            ("Ruta del proyecto", "/var/www/portal-derecho o ruta asignada"),
            ("Permisos", "Comandos autorizados y política de sudo"),
        ],
        [2300, 7060],
    )

    add_heading(doc, "Requisitos del espacio de aplicación", 2)
    add_body(
        doc,
        "El espacio asignado debe permitir desplegar Laravel y ejecutar comandos de Composer y Artisan. Si el usuario no dispone de privilegios de administración, TI deberá instalar y mantener PHP, sus extensiones y el servidor web.",
    )
    for item in (
        "PHP 8.4.1 o superior compatible, preferentemente la última revisión estable de PHP 8.4 validada por TI.",
        "Extensiones pdo_pgsql, pgsql, gd, intl, zip, bcmath, exif, opcache y pcntl.",
        "Composer 2 para instalar dependencias con el archivo composer.lock.",
        "Permiso de escritura en storage y bootstrap/cache.",
        "Raíz pública del sitio configurada en el directorio public del proyecto.",
        "Límite de carga suficiente para PDF de hasta 20 MB e imágenes de hasta 5 MB; se recomienda permitir solicitudes de 40 MB.",
        "Servicio disponible permanentemente, sin suspensión automática por inactividad.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Base de datos PostgreSQL", 2)
    add_body(
        doc,
        "La base de datos puede encontrarse en el mismo servidor o en otro servidor institucional. No es necesario otorgar acceso administrativo al sistema PostgreSQL; basta con crear una base de datos y un usuario exclusivo para la aplicación, con permisos para ejecutar las migraciones y operar las tablas del portal.",
    )
    add_body(
        doc,
        "La conexión debe ser privada y permitirse desde el servidor de la aplicación. PostgreSQL no debe exponerse directamente a Internet. TI debe informar el host, puerto, nombre de la base, usuario, contraseña y modo SSL aplicable. Para la etapa inicial se recomienda reservar 10 GB ampliables y respaldar la base diariamente.",
    )
    add_table_label(doc, 3, "Datos requeridos para PostgreSQL")
    add_apa_table(
        doc,
        ["Parámetro", "Información requerida"],
        [
            ("DB_HOST", "Nombre o IP privada del servidor PostgreSQL"),
            ("DB_PORT", "Puerto, normalmente 5432"),
            ("DB_DATABASE", "Nombre de la base asignada al portal"),
            ("DB_USERNAME", "Usuario exclusivo de la aplicación"),
            ("DB_PASSWORD", "Contraseña entregada mediante canal seguro"),
            ("DB_SSLMODE", "Modo definido por TI: require, verify-full u otro"),
            ("Respaldo", "Frecuencia, retención y responsable de restauración"),
        ],
        [2300, 7060],
    )

    add_heading(doc, "Almacenamiento de documentos e imágenes", 2)
    add_body(
        doc,
        "Los archivos cargados desde el panel administrativo no se guardan dentro de PostgreSQL. La base de datos conserva sus referencias y metadatos, mientras que los PDF, fotografías, portadas e imágenes se almacenan en un directorio o servicio de archivos. Por esta razón, el almacenamiento debe ser persistente y conservarse entre reinicios y despliegues.",
    )
    add_body(
        doc,
        "Si la Universidad utiliza almacenamiento en disco, TI debe entregar una ruta con permisos de lectura y escritura para la aplicación. Se solicita un mínimo inicial de 20 GB y se recomiendan 50 GB ampliables. El directorio debe incluirse en las copias de seguridad. No debe utilizarse un directorio temporal o efímero.",
    )
    add_body(
        doc,
        "Si TI utiliza un proveedor de almacenamiento de objetos, no debe proporcionar otra instancia de servidor. Debe entregar el nombre del bucket, endpoint, región, credenciales, permisos, URL y política de acceso. Esta modalidad requerirá configurar el adaptador correspondiente y migrar los archivos existentes antes de publicar el portal.",
    )
    add_table_label(doc, 4, "Alternativas de almacenamiento")
    add_apa_table(
        doc,
        ["Alternativa", "Entrega requerida", "Adecuación"],
        [
            ("Disco persistente", "Ruta, cuota, permisos y respaldo", "Configuración y enlace del directorio público"),
            ("Objetos tipo S3", "Bucket, endpoint, región, claves y URL", "Configuración adicional y migración de medios"),
        ],
        [2000, 3700, 3660],
    )

    add_heading(doc, "Subdominio, HTTPS y acceso público", 2)
    add_body(
        doc,
        "El servidor puede permanecer dentro de la red privada, pero el portal debe publicarse mediante el mecanismo institucional que defina TI, como proxy inverso, firewall o balanceador. El subdominio propuesto es derecho.unasam.edu.pe, sujeto a disponibilidad y aprobación. TI debe crear el registro DNS, instalar o gestionar el certificado TLS/SSL y redirigir las solicitudes HTTP hacia HTTPS.",
    )
    add_body(
        doc,
        "La conexión SSH puede restringirse a la red de la UNASAM o a VPN sin impedir que el sitio web sea público. Son servicios distintos: SSH se utiliza para administración técnica y HTTPS para que los visitantes consulten el portal.",
    )

    add_heading(doc, "Respaldos, operación y ambiente", 2)
    for item in (
        "Respaldar diariamente la base de datos PostgreSQL y el directorio o bucket de archivos.",
        "Conservar, como referencia inicial, 30 días de copias y cierres mensuales según la política institucional.",
        "Designar al responsable que atenderá restauraciones, cambios de DNS, certificados y fallos del servidor.",
        "Facilitar acceso a los registros de aplicación o un mecanismo para solicitarlos a TI.",
        "Confirmar si el servidor denominado TESTER es únicamente de pruebas o está autorizado para producción.",
        "De ser posible, mantener un ambiente de pruebas separado antes de publicar cambios en el dominio definitivo.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Adecuaciones necesarias en el sistema", 2)
    add_body(
        doc,
        "Cuando TI entregue los accesos será necesario configurar el entorno de producción. El cambio de Neon a PostgreSQL institucional no exige reescribir la aplicación; se deben sustituir las variables de conexión, respaldar y migrar los datos, ejecutar las migraciones y validar el resultado.",
    )
    add_body(
        doc,
        "También se configurarán el dominio definitivo, el certificado a través del proxy institucional, las direcciones de proxies confiables, las cookies seguras, los permisos del almacenamiento y la publicación de archivos. Si TI proporciona almacenamiento tipo S3 o varias instancias de aplicación, serán necesarias adecuaciones adicionales para el sistema de archivos, caché y sesiones compartidas.",
    )

    add_heading(doc, "Secuencia propuesta de implementación", 2)
    for item in (
        "TI confirma el tipo de alojamiento y entrega los datos de acceso mediante un canal seguro.",
        "El responsable técnico valida SSH, versión de PHP, extensiones, permisos y conectividad con PostgreSQL.",
        "Se realiza una copia de seguridad de Neon y de los archivos actuales.",
        "Se despliega el código, se configura el entorno y se migran la base de datos y los medios.",
        "Se prueban el sitio público, panel administrativo, carga de archivos, permisos, vistas previas y ruta de salud.",
        "TI configura el subdominio y HTTPS; la Facultad aprueba la publicación definitiva.",
        "Se documentan los accesos, respaldos, responsables, mantenimiento y procedimiento de recuperación.",
    ):
        add_number(doc, item)

    add_heading(doc, "Texto sugerido para el pedido institucional", 2)
    add_body(
        doc,
        "Asunto: Solicitud de espacio de alojamiento institucional y servicios técnicos para el Portal Web de la Facultad de Derecho y Ciencias Políticas.",
        indent=False,
        bold_lead="Asunto:",
    )
    add_body(
        doc,
        "Se solicita al Área de Tecnologías de la Información la asignación de un espacio de producción en un servidor Linux institucional para alojar el Portal Web de la Facultad de Derecho y Ciencias Políticas. El espacio podrá formar parte de un servidor compartido, máquina virtual o plataforma institucional equivalente, siempre que permita ejecutar una aplicación desarrollada con Laravel 13.12.0, Filament 5.6.6 y PHP 8.4.1 o superior compatible.",
    )
    add_body(
        doc,
        "Se requiere acceso técnico mediante SSH restringido a la red institucional o mediante VPN, proporcionando nombre o IP del servidor, puerto, usuario, mecanismo de autenticación, ruta asignada y permisos disponibles. Asimismo, se solicita una base de datos PostgreSQL institucional con usuario exclusivo, un directorio persistente de al menos 20 GB para documentos e imágenes, copias de seguridad y un responsable técnico de coordinación.",
    )
    add_body(
        doc,
        "Finalmente, se solicita la creación y publicación del subdominio institucional que se apruebe, proponiéndose derecho.unasam.edu.pe, junto con su registro DNS, certificado TLS/SSL y mecanismo de publicación HTTPS. Las credenciales deberán entregarse mediante un canal seguro y no deberán incorporarse en el oficio ni en el repositorio del sistema.",
    )

    add_heading(doc, "Conclusión", 2)
    add_body(
        doc,
        "La entrega puede realizarse mediante un documento simple con servidor, IP, puerto, usuario y contraseña o llave, como ocurrió con TESTER. Debe incluir también los datos de PostgreSQL, la ubicación persistente de archivos, el respaldo, el subdominio y el responsable operativo. Esta información permitirá configurar, migrar y publicar el sistema de manera controlada.",
    )


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    properties = doc.core_properties
    properties.title = "Informe técnico para la solicitud de alojamiento institucional del Portal Web de Derecho"
    properties.subject = "Acceso SSH, PostgreSQL, almacenamiento persistente y subdominio institucional"
    properties.author = "Facultad de Derecho y Ciencias Políticas - UNASAM"
    properties.keywords = "UNASAM, SSH, Laravel, PostgreSQL, alojamiento institucional"
    properties.comments = "Documento formal con presentación basada en APA 7"

    add_title_page(doc)
    add_summary(doc)
    add_main_content(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
