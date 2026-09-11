from __future__ import annotations

from datetime import date
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "infrastructure" / "Informe_tecnico_requerimientos_despliegue_portal_Derecho_UNASAM.docx"
LOGO = ROOT / "public" / "img" / "escudo-unasam.png"
SKILL_SCRIPTS = Path(
    "/Users/yomeljairbarretoflores/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.805.11740/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = RGBColor(18, 42, 76)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GOLD = RGBColor(178, 137, 61)
INK = RGBColor(32, 38, 46)
GRAY = RGBColor(93, 103, 115)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
PALE_GOLD = "FFF6DF"
WHITE = RGBColor(255, 255, 255)
PAGE_WIDTH_DXA = 9360


def set_run_font(run, *, size=11, bold=False, italic=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, *, color="D8DEE8", size=6):
    color = str(color)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_paragraph(doc, text="", *, bold=False, italic=False, color=INK, size=11,
                  before=0, after=6, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT,
                  keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_rich_paragraph(doc, parts, *, before=0, after=6, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, options in parts:
        set_run_font(p.add_run(text), **options)
    return p


def add_bullet(doc, text, *, level=0, after=5):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text), size=11)
    return p


def add_number(doc, text, *, after=5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text), size=11)
    return p


def add_callout(doc, title, text, *, fill=PALE_GOLD, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(title + " "), size=11, bold=True, color=NAVY)
    set_run_font(p.add_run(text), size=11, color=INK)
    apply_table_geometry(table, [PAGE_WIDTH_DXA], indent_dxa=180,
                         cell_margins_dxa={"top": 140, "bottom": 140, "start": 180, "end": 180})
    add_paragraph(doc, after=3)
    return table


def add_table(doc, headers, rows, widths, *, header_fill=LIGHT_BLUE, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, heading in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        set_run_font(p.add_run(heading), size=font_size, bold=True, color=NAVY)
    for values in rows:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            set_run_font(p.add_run(str(value)), size=font_size, color=INK)
    apply_table_geometry(
        table,
        widths,
        indent_dxa=130,
        cell_margins_dxa={"top": 100, "bottom": 100, "start": 130, "end": 130},
    )
    add_paragraph(doc, after=2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def configure_numbering_styles(doc):
    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(5)
    bullet.paragraph_format.line_spacing = 1.167

    bullet2 = doc.styles["List Bullet 2"]
    bullet2.font.name = "Calibri"
    bullet2.font.size = Pt(11)
    bullet2.paragraph_format.left_indent = Inches(0.75)
    bullet2.paragraph_format.first_line_indent = Inches(-0.25)
    bullet2.paragraph_format.space_after = Pt(4)
    bullet2.paragraph_format.line_spacing = 1.167

    numbered = doc.styles["List Number"]
    numbered.font.name = "Calibri"
    numbered.font.size = Pt(11)
    numbered.paragraph_format.left_indent = Inches(0.5)
    numbered.paragraph_format.first_line_indent = Inches(-0.25)
    numbered.paragraph_format.space_after = Pt(5)
    numbered.paragraph_format.line_spacing = 1.167


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    values = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in values.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    configure_numbering_styles(doc)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("Página "), size=8.5, color=GRAY)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=8.5, color=GRAY)


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("UNASAM  |  FACULTAD DE DERECHO Y CIENCIAS POLÍTICAS"), size=8.5, bold=True, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    add_page_number(fp)


def add_cover(doc):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(26)
        p.paragraph_format.space_after = Pt(14)
        logo = p.add_run().add_picture(str(LOGO), width=Inches(0.86))
        logo._inline.docPr.set("descr", "Escudo institucional de la UNASAM")
        logo._inline.docPr.set("title", "Escudo de la Universidad Nacional Santiago Antúnez de Mayolo")

    add_paragraph(doc, "UNIVERSIDAD NACIONAL SANTIAGO ANTÚNEZ DE MAYOLO",
                  size=10, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_paragraph(doc, "Facultad de Derecho y Ciencias Políticas",
                  size=13, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    add_paragraph(doc, "INFORME TÉCNICO",
                  size=11, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)
    add_paragraph(doc, "Requerimientos de infraestructura y servicios",
                  size=26, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=5, line=1.0)
    add_paragraph(doc, "para el despliegue institucional del Portal Web",
                  size=18, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=22, line=1.0)
    add_paragraph(doc, "Documento de especificación para solicitud al área de Tecnologías de la Información",
                  size=11.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)

    table = doc.add_table(rows=4, cols=2)
    metadata = [
        ("Sistema", "Portal de Derecho y Ciencias Políticas - UNASAM"),
        ("Dominio propuesto", "derecho.unasam.edu.pe (sujeto a disponibilidad y aprobación)"),
        ("Versión técnica", "Laravel 13.12.0 / PHP 8.4.1+ / PostgreSQL / Filament 5.6.6"),
        ("Fecha", "12 de agosto de 2026"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=9.5, bold=True, color=NAVY)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=9.5, color=INK)
    apply_table_geometry(table, [2700, 6660], indent_dxa=140,
                         cell_margins_dxa={"top": 110, "bottom": 110, "start": 140, "end": 140})

    add_paragraph(doc, "Uso previsto: sustento técnico para el requerimiento de alojamiento, publicación y operación del sistema.",
                  size=9, color=GRAY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0)
    doc.add_page_break()


def add_executive_summary(doc):
    add_heading(doc, "1. Resumen ejecutivo", 1)
    add_paragraph(
        doc,
        "Para publicar el portal en la infraestructura institucional no basta solicitar una dirección web ni una cuenta de acceso. Se requiere que el área de Tecnologías de la Información asigne o provea una plataforma de producción completa, con capacidad de cómputo, base de datos PostgreSQL, almacenamiento persistente para archivos, publicación HTTPS bajo el dominio unasam.edu.pe, copias de seguridad y un mecanismo formal de operación técnica.",
    )
    add_callout(
        doc,
        "Recomendación principal.",
        "Solicitar una máquina virtual Linux exclusiva o un servicio institucional de contenedores con recursos equivalentes. El acceso SSH es apropiado si se entrega una máquina virtual; en una plataforma administrada debe sustituirse por acceso al despliegue, variables de entorno, consola, logs y reinicio del servicio.",
    )
    add_paragraph(doc, "La solicitud debe contemplar, como mínimo:", bold=True, color=NAVY, after=5)
    for item in [
        "Subdominio institucional y certificado TLS/SSL gestionado por TI.",
        "Entorno de ejecución Linux para la aplicación Laravel.",
        "Base de datos PostgreSQL administrada, aislada y respaldada.",
        "Almacenamiento persistente para PDF, imágenes, fotografías y documentos.",
        "Acceso técnico controlado para desplegar, migrar, revisar logs y recuperar el servicio.",
        "Política de copias de seguridad, monitoreo, soporte y responsables operativos.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Alcance y arquitectura del sistema", 1)
    add_paragraph(
        doc,
        "El sistema es un portal institucional con panel administrativo. Publica información académica, estadísticas, documentos normativos, comunicados, noticias, plana docente y la revista científica «Derecho y Cultura», incluyendo números, artículos, portadas y archivos PDF. La aplicación mantiene trazabilidad editorial, perfiles administrativos, control de publicación, caché pública y cabeceras de seguridad.",
    )
    add_table(
        doc,
        ["Capa", "Tecnología actual", "Necesidad en producción"],
        [
            ("Aplicación", "Laravel 13.12.0, PHP 8.4.1+, Filament 5.6.6, Livewire", "Runtime PHP compatible y servicio web permanente"),
            ("Interfaz", "Blade, Tailwind CSS 4, Vite 8", "Assets compilados durante CI o construcción del contenedor"),
            ("Base de datos", "PostgreSQL", "Instancia o base institucional con usuario propio y respaldo"),
            ("Archivos", "Spatie Media Library", "Disco persistente o almacenamiento de objetos compatible"),
            ("Caché y sesión", "Archivos locales en una sola instancia", "Volumen escribible; Redis si existen múltiples réplicas"),
            ("Entrega", "Docker disponible", "Registro/repositorio y procedimiento institucional de despliegue"),
        ],
        [1500, 2800, 5060],
        font_size=9.1,
    )


def add_request_requirements(doc):
    add_heading(doc, "3. Requerimiento que debe formularse al área de TI", 1)
    add_paragraph(doc, "Se recomienda que el oficio o solicitud institucional pida expresamente los siguientes componentes y no utilice únicamente la frase «una instancia de servidor».")

    add_heading(doc, "3.1. Dominio, DNS y publicación HTTPS", 2)
    for item in [
        "Reserva y creación del subdominio derecho.unasam.edu.pe, u otro nombre que apruebe la Universidad.",
        "Registro DNS tipo A/AAAA o CNAME apuntando al servicio definitivo.",
        "Certificado TLS/SSL válido, renovación automática y redirección obligatoria de HTTP a HTTPS.",
        "Definición del proxy inverso, balanceador o firewall institucional que publicará los puertos 80 y 443.",
        "Entrega de las IP o redes CIDR de los proxies institucionales para configurar TRUSTED_PROXIES; no se aceptará el comodín «*».",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.2. Plataforma de aplicación", 2)
    add_paragraph(doc, "Opción recomendada: una máquina virtual Linux exclusiva de producción o un servicio de contenedores equivalente. Requerimientos iniciales:")
    add_table(
        doc,
        ["Recurso", "Mínimo inicial", "Recomendado", "Observación"],
        [
            ("CPU", "2 vCPU", "4 vCPU", "El portal es liviano; las conversiones de imágenes generan picos"),
            ("Memoria RAM", "4 GB", "8 GB", "Incluye PHP, servidor web, caché y procesos de despliegue"),
            ("Sistema operativo", "Linux 64 bits", "Ubuntu Server 24.04 LTS o equivalente", "Mantenimiento y parches a cargo de TI"),
            ("Disco del sistema", "30 GB SSD", "50 GB SSD", "Código, dependencias, logs temporales y espacio de despliegue"),
            ("Disponibilidad", "Servicio continuo", "Monitoreo 24x7 y reinicio automático", "Evitar suspensión por inactividad"),
        ],
        [1500, 1450, 2100, 4310],
        font_size=8.9,
    )
    add_paragraph(doc, "El proyecto ya dispone de Dockerfile basado en PHP 8.4. Si TI acepta contenedores, puede desplegarse con una imagen construida desde el repositorio y un volumen persistente. Si TI no utiliza Docker, deberá habilitar PHP 8.4.1 o superior dentro de la rama 8.4/8.5, con las extensiones indicadas en el Anexo A, Composer 2 y un servidor web compatible.")

    add_heading(doc, "3.3. Base de datos PostgreSQL", 2)
    for item in [
        "PostgreSQL 15 o superior; se recomienda una versión con soporte institucional vigente.",
        "Base de datos exclusiva para producción y usuario técnico exclusivo de la aplicación, sin privilegios de superusuario.",
        "Host, puerto, nombre de base, usuario, contraseña y modo SSL/TLS. Las credenciales deben entregarse por un canal seguro y no incluirse en el oficio ni en el repositorio.",
        "Conectividad permitida únicamente desde el servidor de aplicación o su red autorizada; PostgreSQL no debe exponerse públicamente.",
        "Zona horaria America/Lima y codificación UTF-8.",
        "Copias de seguridad automáticas, cifradas y con pruebas periódicas de restauración.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "Dimensionamiento inicial.", "Reservar 10 GB para la base de datos es suficiente para la fase inicial y deja margen amplio para contenido, auditoría y crecimiento. Debe existir capacidad de ampliación sin reinstalar el sistema.", fill=LIGHT_BLUE, accent="8CA7C5")

    add_heading(doc, "3.4. Almacenamiento de archivos", 2)
    add_paragraph(doc, "Este componente es obligatorio. La base de datos conserva metadatos, pero los PDF, fotografías, portadas e imágenes se almacenan fuera de PostgreSQL. Hay dos alternativas válidas:")
    add_table(
        doc,
        ["Alternativa", "Qué debe proporcionar TI", "Implicación técnica"],
        [
            ("Disco persistente", "Volumen SSD montado en el servidor, respaldado y conservado entre despliegues", "Es la alternativa más simple para una sola instancia; se usa storage/app/public y public/storage"),
            ("Almacenamiento de objetos", "Bucket privado o público controlado, endpoint, región, credenciales y política CORS/URL", "Requiere configurar el disco S3 compatible y migrar los medios existentes"),
        ],
        [1700, 3600, 4060],
        font_size=9.1,
    )
    add_paragraph(doc, "Capacidad solicitada:", bold=True, color=NAVY, after=4)
    for item in [
        "Mínimo inicial: 20 GB persistentes para medios; recomendado: 50 GB con posibilidad de ampliación.",
        "La aplicación admite actualmente imágenes de hasta 5 MB y PDF de hasta 20 MB por archivo.",
        "El servidor web, PHP y cualquier proxy deben aceptar al menos 25 MB por archivo y 30 MB por petición; se recomienda 32 MB y 40 MB respectivamente para evitar rechazos por sobrecarga del formulario.",
        "Debe respaldarse tanto el archivo original como sus conversiones/miniaturas y mantenerse sincronía con la tabla media de PostgreSQL.",
        "No debe utilizarse un disco efímero. Si un despliegue elimina el volumen, se perderán artículos, resoluciones, fotografías y portadas.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.5. Accesos técnicos y seguridad operativa", 2)
    add_paragraph(doc, "El acceso requerido depende del modelo de alojamiento:")
    add_table(
        doc,
        ["Modelo", "Acceso solicitado"],
        [
            ("Máquina virtual", "Cuenta SSH nominativa mediante llave pública, privilegios sudo limitados, acceso SFTP al volumen y permiso para ejecutar despliegues"),
            ("Contenedores / PaaS", "Acceso al proyecto, registro de imágenes, variables/secretos, consola de ejecución, logs, métricas, reinicio y rollback"),
            ("Hosting administrado", "Panel de despliegue, variables de entorno, cron si corresponde, logs y canal de soporte; debe garantizar PHP/PostgreSQL/almacenamiento compatibles"),
        ],
        [2300, 7060],
        font_size=9.4,
    )
    for item in [
        "Repositorio Git institucional o acceso controlado al repositorio autorizado para desplegar la rama de producción.",
        "Separación entre credenciales del sistema operativo, base de datos, panel administrativo y almacenamiento.",
        "Gestión de secretos fuera del código: APP_KEY, credenciales PostgreSQL, correo y almacenamiento.",
        "Acceso a logs de aplicación y servidor sin exponer contraseñas ni datos sensibles.",
        "Ambiente de pruebas o preproducción recomendado para validar migraciones antes del dominio público.",
    ]:
        add_bullet(doc, item)


def add_operations(doc):
    add_heading(doc, "4. Servicios complementarios y operación", 1)

    add_heading(doc, "4.1. Correo institucional", 2)
    add_paragraph(doc, "El portal funciona actualmente sin recuperación automática de contraseña. Para habilitar notificaciones o recuperación por correo se debe solicitar una cuenta o relay SMTP institucional, por ejemplo no-responder@unasam.edu.pe o una dirección aprobada por TI, con host, puerto, cifrado, usuario, contraseña y restricciones de envío. Mientras no exista SMTP, la recuperación seguirá siendo administrada desde el panel por un superadministrador.")

    add_heading(doc, "4.2. Copias de seguridad y recuperación", 2)
    add_paragraph(doc, "La copia de seguridad debe tratar la base de datos y los archivos como una sola unidad lógica. Se propone la siguiente política mínima:")
    add_table(
        doc,
        ["Activo", "Frecuencia", "Retención mínima", "Prueba"],
        [
            ("PostgreSQL", "Diaria; incremental/continua si TI dispone", "30 días y 12 cierres mensuales", "Restauración trimestral"),
            ("Archivos y medios", "Diaria", "30 días y 12 cierres mensuales", "Recuperación de muestra trimestral"),
            ("Configuración y secretos", "Ante cada cambio", "Versiones vigentes y anterior", "Validación semestral"),
            ("Código fuente", "Por cada cambio en Git", "Historial permanente", "Despliegue de rollback"),
        ],
        [2200, 2200, 2800, 2160],
        font_size=9.1,
    )
    add_paragraph(doc, "Objetivos sugeridos: RPO máximo de 24 horas y RTO inicial de 8 horas. TI puede establecer valores más exigentes conforme a su política de continuidad.")

    add_heading(doc, "4.3. Monitoreo, registros y mantenimiento", 2)
    for item in [
        "Monitoreo HTTP del health check /up y del certificado TLS.",
        "Alertas por indisponibilidad, errores 5xx, falta de espacio y fallos de respaldo.",
        "Retención de logs de aplicación y servidor por 30 días como mínimo, con rotación.",
        "Actualizaciones de seguridad del sistema operativo, PHP, PostgreSQL y dependencias coordinadas entre TI y el responsable técnico.",
        "Procedimiento de mantenimiento, rollback y contacto de escalamiento.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. Configuraciones y adecuaciones requeridas en el código", 1)
    add_paragraph(doc, "Sí, cuando TI entregue la infraestructura será necesario adaptar la configuración de producción. La mayor parte se realiza mediante variables de entorno y despliegue; solo el tipo de almacenamiento puede requerir cambios de código adicionales.")
    add_table(
        doc,
        ["Componente", "Configuración prevista", "¿Cambio de código?"],
        [
            ("Dominio/HTTPS", "APP_URL, cookie segura, DNS, proxy y TRUSTED_PROXIES", "Normalmente no"),
            ("PostgreSQL", "DB_HOST, DB_PORT, DB_DATABASE, DB_USERNAME, DB_PASSWORD y DB_SSLMODE", "No, salvo requisitos particulares de conexión"),
            ("Disco persistente", "Montaje, permisos, FILESYSTEM_DISK/MEDIA_DISK y storage:link", "No o mínimo"),
            ("Objeto S3 compatible", "Bucket, endpoint, región, credenciales y URLs", "Sí: validar/configurar Media Library y migrar medios"),
            ("Correo SMTP", "MAIL_MAILER, host, puerto, cifrado y remitente", "Mínimo; recuperación requiere habilitar su flujo"),
            ("Múltiples instancias", "Redis para caché/sesiones y almacenamiento compartido", "Sí, configuración de arquitectura"),
            ("Servidor sin Docker", "Nginx/Apache/FrankenPHP, PHP-FPM, extensiones y límites", "No en lógica; sí en automatización"),
        ],
        [1900, 4900, 2560],
        font_size=8.9,
    )
    add_paragraph(doc, "Además, durante la migración se deberá:", bold=True, color=NAVY, after=4)
    for item in [
        "Generar una APP_KEY exclusiva de producción y preservar esa clave durante toda la vida del sistema.",
        "Migrar el esquema y los datos desde Neon hacia PostgreSQL institucional mediante un respaldo controlado.",
        "Copiar los archivos existentes al volumen o bucket definitivo y verificar las referencias de Media Library.",
        "Ejecutar migraciones, storage:link, optimización de configuración/rutas/vistas y calentamiento de caché.",
        "Crear o validar cuentas administrativas sin trasladar contraseñas por canales inseguros.",
        "Probar rutas públicas, panel, carga/descarga de archivos, permisos, sitemap, vistas previas y health check.",
    ]:
        add_bullet(doc, item)


def add_deployment_plan(doc):
    add_heading(doc, "6. Secuencia recomendada de implementación", 1)
    steps = [
        "TI confirma el modelo de alojamiento, recursos, responsable, subdominio y fecha de intervención.",
        "TI entrega accesos temporales y secretos mediante un canal seguro; el técnico valida conectividad y permisos.",
        "Se crea el entorno de preproducción y se despliega la aplicación con PostgreSQL y almacenamiento definitivos.",
        "Se respalda Neon y los medios actuales; luego se migran datos y archivos al entorno institucional.",
        "Se ejecutan migraciones y pruebas funcionales, de seguridad, rendimiento y recuperación.",
        "TI configura DNS y certificado; se realiza una ventana de corte para publicación definitiva.",
        "Se valida el portal público y el panel; se documenta el rollback y se entrega la operación.",
        "Se revocan accesos temporales, se rotan credenciales y se inicia la política de respaldo y monitoreo.",
    ]
    for step in steps:
        add_number(doc, step)

    add_heading(doc, "7. Responsabilidades propuestas", 1)
    add_table(
        doc,
        ["Responsable", "Funciones principales"],
        [
            ("Área de TI de la UNASAM", "Infraestructura, red, DNS, TLS, PostgreSQL, almacenamiento, respaldos, monitoreo, parches y custodia de secretos"),
            ("Facultad de Derecho y Ciencias Políticas", "Aprobación del dominio, designación de administradores, propiedad y validación del contenido institucional"),
            ("Responsable técnico del portal", "Configuración de la aplicación, migración, despliegue, pruebas, documentación y soporte de la lógica del sistema"),
            ("Editores autorizados", "Carga, revisión, publicación y actualización del contenido desde el panel administrativo"),
        ],
        [2850, 6510],
        font_size=9.2,
    )

    add_heading(doc, "8. Criterios de aceptación", 1)
    for item in [
        "El subdominio institucional carga exclusivamente mediante HTTPS y redirige HTTP.",
        "El health check /up responde satisfactoriamente y el servicio no se suspende por inactividad.",
        "La aplicación se conecta a PostgreSQL institucional sin exposición pública de la base.",
        "El panel permite subir, visualizar, descargar y eliminar de forma controlada PDF e imágenes persistentes.",
        "Un nuevo despliegue o reinicio no elimina los archivos cargados.",
        "Las copias de seguridad de base y archivos se ejecutan y existe evidencia de restauración.",
        "Los logs y métricas son accesibles para los responsables autorizados.",
        "Se comprueban panel, permisos, vistas previas, contenido público, sitemap, seguridad y rendimiento.",
    ]:
        add_bullet(doc, item)


def add_request_text(doc):
    doc.add_page_break()
    add_heading(doc, "9. Texto sugerido para el pedido institucional", 1)
    add_callout(
        doc,
        "Asunto sugerido:",
        "Solicitud de infraestructura y servicios tecnológicos para el despliegue del Portal Web de la Facultad de Derecho y Ciencias Políticas.",
        fill=LIGHT_BLUE,
        accent="8CA7C5",
    )
    add_paragraph(
        doc,
        "Se solicita al área competente la asignación y configuración de una plataforma de producción para el Portal Web de la Facultad de Derecho y Ciencias Políticas, desarrollado con Laravel 13.12.0, Filament 5.6.6, PHP 8.4.1 o superior y PostgreSQL. La plataforma deberá contemplar un entorno Linux o servicio de contenedores equivalente, base de datos PostgreSQL exclusiva y respaldada, almacenamiento persistente para documentos e imágenes, publicación bajo un subdominio institucional (propuesto: derecho.unasam.edu.pe), certificado TLS/SSL, monitoreo, registros y copias de seguridad.",
        italic=True,
    )
    add_paragraph(
        doc,
        "Asimismo, se solicita designar un responsable técnico y otorgar al personal autorizado los accesos necesarios para el despliegue y mantenimiento: cuenta SSH con llave y privilegios limitados cuando se trate de una máquina virtual, o acceso equivalente a consola, variables de entorno, logs, reinicio y rollback cuando se trate de una plataforma administrada. Las credenciales y secretos deberán ser entregados mediante un canal seguro.",
        italic=True,
    )
    add_paragraph(
        doc,
        "Como dimensionamiento inicial se propone 2 vCPU, 4 GB de RAM, 30 GB de disco para el sistema, 10 GB para PostgreSQL y 20 GB de almacenamiento persistente para archivos, recomendándose 4 vCPU, 8 GB de RAM, 50 GB de sistema y 50 GB de medios, con posibilidad de ampliación. Se adjunta el informe técnico con las especificaciones, alternativas de almacenamiento, política de respaldo y criterios de aceptación.",
        italic=True,
    )


def add_appendices(doc):
    add_heading(doc, "Anexo A. Requisitos técnicos de compatibilidad", 1)
    add_table(
        doc,
        ["Elemento", "Requisito"],
        [
            ("PHP", "8.4.1 como mínimo efectivo; se recomienda la última revisión estable de PHP 8.4 compatible con el proyecto"),
            ("Extensiones PHP", "pdo_pgsql, pgsql, gd, intl, zip, bcmath, exif, opcache y pcntl"),
            ("Composer", "Versión 2, instalación de dependencias con --no-dev y autoload optimizado"),
            ("Node.js", "22, necesario en CI/build para npm ci y Vite; no es obligatorio en ejecución si los assets ya están compilados"),
            ("PostgreSQL", "15 o superior recomendado, UTF-8, conexión privada y TLS según política institucional"),
            ("Servidor web", "FrankenPHP/Caddy, Nginx + PHP-FPM o Apache compatible; document root en public/"),
            ("Escritura", "storage/ y bootstrap/cache; volumen persistente para storage/app/public"),
            ("Procesos", "No se requieren workers ni scheduler dedicados en la configuración actual; QUEUE_CONNECTION=sync"),
            ("Salud", "Ruta /up"),
        ],
        [2200, 7160],
        font_size=9.2,
    )

    add_heading(doc, "Anexo B. Variables y datos que TI debe entregar", 1)
    checklist = [
        ("Dominio", "Subdominio aprobado, destino DNS y fecha de propagación"),
        ("Proxy", "IPs/CIDR autorizadas y cabeceras X-Forwarded-*"),
        ("Servidor", "Host/IP privada, método de acceso, usuario, puerto y política de sudo"),
        ("PostgreSQL", "Host, puerto, base, usuario, contraseña, SSLMODE y responsable de respaldos"),
        ("Archivos", "Ruta del volumen o datos del bucket, capacidad, permisos, URL y respaldo"),
        ("Correo", "Host SMTP, puerto, cifrado, usuario, contraseña y remitente aprobado, si se habilita"),
        ("Logs", "Ruta/panel de consulta, retención y canal de alertas"),
        ("Operación", "Responsable TI, ventana de mantenimiento y procedimiento de escalamiento"),
    ]
    add_table(doc, ["Dato", "Información requerida"], checklist, [1800, 7560], font_size=9.1)

    add_heading(doc, "Anexo C. Variables de producción principales", 1)
    add_table(
        doc,
        ["Variable", "Finalidad / valor esperado"],
        [
            ("APP_ENV", "production"),
            ("APP_DEBUG", "false"),
            ("APP_URL", "https://derecho.unasam.edu.pe o dominio aprobado"),
            ("APP_KEY", "Clave exclusiva generada para producción; no debe regenerarse"),
            ("TRUSTED_PROXIES", "IPs o CIDR explícitos entregados por TI"),
            ("DB_CONNECTION", "pgsql"),
            ("DB_*", "Parámetros de PostgreSQL institucional"),
            ("DB_SSLMODE", "require, verify-full o el valor definido por TI"),
            ("MEDIA_UPLOADS_ENABLED", "true cuando el almacenamiento persistente esté validado"),
            ("FILESYSTEM_DISK / MEDIA_DISK", "public para disco local; s3 para proveedor compatible"),
            ("SESSION_SECURE_COOKIE", "true"),
            ("SESSION_DRIVER / CACHE_STORE", "file en una instancia; redis en múltiples instancias"),
            ("LOG_CHANNEL", "stderr, syslog o canal institucional"),
            ("MAIL_*", "Solo cuando TI provea SMTP"),
        ],
        [2700, 6660],
        font_size=8.9,
    )

    add_heading(doc, "Conclusión", 1)
    add_paragraph(
        doc,
        "La Facultad debe solicitar una solución de alojamiento institucional integral y no solamente una cuenta SSH. La cuenta SSH es un medio de administración posible, pero el requerimiento real comprende cómputo, PostgreSQL, almacenamiento persistente, dominio HTTPS, accesos operativos, respaldos y soporte. Una vez que TI defina y entregue estos componentes, el responsable técnico podrá configurar el proyecto, migrar la información desde Neon, trasladar los archivos y ejecutar las pruebas de puesta en producción.",
        bold=True,
        color=NAVY,
        after=12,
    )
    add_callout(
        doc,
        "Decisión solicitada a TI.",
        "Confirmar el modelo de alojamiento, aprobar el subdominio, reservar los recursos y designar un responsable técnico institucional para coordinar la puesta en producción.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_heading(doc, "Siguientes acciones", 2)
    for index, action in enumerate(
        [
            "La Facultad remite el pedido institucional acompañado de este informe.",
            "TI comunica el modelo disponible y completa los datos del Anexo B.",
            "El responsable técnico valida accesos, prepara preproducción y presenta el plan de migración.",
            "La Facultad y TI aprueban la ventana de publicación y el esquema de operación posterior.",
        ],
        start=1,
    ):
        paragraph = add_paragraph(doc, f"{index}.  {action}", after=5)
        paragraph.paragraph_format.left_indent = Twips(360)
        paragraph.paragraph_format.first_line_indent = Twips(-180)

    add_heading(doc, "Conformidad y coordinación", 2)
    add_table(
        doc,
        ["Rol", "Nombre / área", "Fecha y conformidad"],
        [
            ("Facultad solicitante", "", ""),
            ("Área de Tecnologías de la Información", "", ""),
            ("Responsable técnico del portal", "", ""),
        ],
        [2600, 3600, 3160],
        font_size=9.2,
    )


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    props = doc.core_properties
    props.title = "Informe técnico de requerimientos para el despliegue del Portal de Derecho UNASAM"
    props.subject = "Infraestructura, PostgreSQL, almacenamiento, dominio y operación"
    props.author = "Facultad de Derecho y Ciencias Políticas - UNASAM"
    props.keywords = "UNASAM, Laravel, PostgreSQL, infraestructura, despliegue, portal web"
    props.comments = "Documento técnico para solicitud institucional"

    add_cover(doc)
    add_executive_summary(doc)
    add_request_requirements(doc)
    add_operations(doc)
    add_deployment_plan(doc)
    add_request_text(doc)
    add_appendices(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
